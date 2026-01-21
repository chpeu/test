from __future__ import annotations

import json
import argparse
import os
import re
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
import psycopg2
import openpyxl
from psycopg2.extras import RealDictCursor
from urllib.parse import quote_plus


def _num(v: Any) -> Optional[float]:
    try:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        return float(v)
    except Exception:
        return None


def _str(v: Any) -> Optional[str]:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    s = str(v).strip()
    return s if s else None


def _to_utc_ts(v: Any) -> Optional[float]:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return None
    if isinstance(v, (int, float)):
        # Assume ms if very large
        if v > 1e12:
            return float(v) / 1000.0
        if v > 1e10:
            return float(v)
        return float(v)
    if isinstance(v, datetime):
        dt = v
        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)
        return dt.timestamp()
    s = str(v).strip()
    if not s:
        return None
    # pandas handles many formats
    try:
        dt = pd.to_datetime(s, utc=True, errors="coerce")
        if pd.isna(dt):
            return None
        return dt.to_pydatetime().timestamp()
    except Exception:
        return None


def _normalize_symbol(s: Optional[str]) -> Optional[str]:
    if not s:
        return None
    s = s.strip().upper()
    s = _clean_text(s) or s
    s = s.replace("\n", " ").replace("\r", " ")
    s = re.sub(r"\s+", " ", s).strip()

    # Strip contract type suffixes from exchange exports
    s = re.sub(r"\s*(PERPÉTUEL|PERPETUEL|PERPETUAL|PERP)\s*$", "", s, flags=re.IGNORECASE).strip()

    # Remove any remaining spaces/newlines (some exports break long symbols)
    s = re.sub(r"\s+", "", s)

    # Normalize DB formats like 'SUI/USDT:USDT' -> 'SUI_USDT'
    if ":" in s:
        s = s.split(":", 1)[0]
    s = s.replace("/", "_").replace("-", "_")

    # Standardize 'BASEUSDT' -> 'BASE_USDT'
    if s.endswith("USDT") and not s.endswith("_USDT"):
        s = s[: -4] + "_USDT"

    s = re.sub(r"_+", "_", s)
    return s


def _normalize_side_to_direction(s: Optional[str]) -> Optional[str]:
    if not s:
        return None
    v = s.strip().upper()
    mapping = {
        "LONG": "LONG",
        "SHORT": "SHORT",
        "BUY": "LONG",
        "SELL": "SHORT",
        "OPEN_LONG": "LONG",
        "OPEN_SHORT": "SHORT",
        "ACHETER": "LONG",
        "VENDRE": "SHORT",
    }
    return mapping.get(v, v)


def _pick_col(df: pd.DataFrame, candidates: List[str]) -> Optional[str]:
    cols = {c.lower(): c for c in df.columns}
    for cand in candidates:
        if cand.lower() in cols:
            return cols[cand.lower()]
    return None


def _clean_text(s: Any) -> Optional[str]:
    v = _str(s)
    if not v:
        return None
    # remove directional marks
    v = v.replace("\u200e", "").replace("\u200f", "")
    v = v.replace("\xa0", " ")
    return v.strip()


def _parse_qty_to_float(qty: Optional[str]) -> Optional[float]:
    if not qty:
        return None
    s = _clean_text(qty) or ""
    # Keep only leading numeric part, allow commas as thousand separators
    m = re.search(r"([0-9][0-9,\. ]*)", s)
    if not m:
        return None
    num = m.group(1).replace(" ", "").replace(",", "")
    try:
        return float(num)
    except Exception:
        return None


def _parse_usdt_amount(s: Optional[str]) -> Optional[float]:
    if not s:
        return None
    v = _clean_text(s) or ""
    m = re.search(r"([-+]?\d+(?:\.\d+)?)\s*USDT", v.upper())
    if not m:
        return None
    try:
        return float(m.group(1))
    except Exception:
        return None


def _parse_all_usdt_amounts(s: Optional[str]) -> List[float]:
    if not s:
        return []
    v = (_clean_text(s) or "").upper()
    vals = []
    for m in re.finditer(r"([-+]?\d+(?:\.\d+)?)\s*USDT", v):
        try:
            vals.append(float(m.group(1)))
        except Exception:
            continue
    return vals


def _get_contract_size(symbol: str, specs: Dict[str, Any]) -> float:
    if symbol in specs:
        return float(specs[symbol].get("contract_size", 1.0))
    # Fallback to 1.0 if not found
    return 1.0


def parse_mexc_fr_orders_xlsx(path: str, sheet: Optional[str]) -> pd.DataFrame:
    # Load contract specs cache
    specs = {}
    cache_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "data", "contract_specs_cache.json")
    if os.path.exists(cache_path):
        try:
            with open(cache_path, "r", encoding="utf-8") as f:
                data = json.load(f)
                specs = data.get("specs", {})
        except Exception as e:
            print(f"Warning: Could not load contract specs cache: {e}")

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]

    # Find header row containing "Direction"
    header_row = None
    for r in range(1, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        row_txt = [(_clean_text(v) or "").lower() for v in row_vals]
        if any(v == "direction" for v in row_txt):
            header_row = r
            break
    if header_row is None:
        raise RuntimeError("Impossible de détecter la ligne d'en-tête (Direction) dans l'export MEXC")

    header_vals = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    header_txt = [(_clean_text(v) or "").strip() for v in header_vals]
    header_l = [h.lower() for h in header_txt]

    def _find_col(candidates: List[str]) -> Optional[int]:
        for cand in candidates:
            cand_l = cand.strip().lower()
            for i, h in enumerate(header_l):
                if h == cand_l or (cand_l and cand_l in h):
                    return i
        return None

    symbol_idx = _find_col(["paire de trading", "paire", "contract", "instrument", "symbole", "symbol", "cryptomonnaie", "crypto"])
    time_idx = _find_col(["date", "time", "timestamp", "heure"])
    direction_idx = _find_col(["direction", "side", "type"])
    qty_idx = _find_col(
        [
            "montant exécuté",
            "montant execute",
            "quantité exécutée",
            "quantite executee",
            "montant",
            "qty",
            "quantity",
        ]
    )
    exec_price_idx = _find_col(["prix d'exécution", "prix d'execution", "execution price", "prix exécuté", "prix execute"])
    pnl_idx = _find_col(["pertes et profits (p&l)", "p&l", "pnl", "profit", "pertes et profits"])
    fee_idx = _find_col(["frais", "fee", "fees", "commission"])
    status_idx = _find_col(["statut", "status"])
    reduce_only_idx = _find_col(["ordres reduce-only", "ordre reduce-only", "reduce-only", "reduce only"])

    orders = []
    for r in range(header_row + 1, ws.max_row + 1):
        cells = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        # skip progress rows like "100.00%"
        non_empty = [v for v in cells if v is not None and str(v).strip() != ""]
        if not non_empty:
            continue
        if len(non_empty) == 1 and _clean_text(non_empty[0]) and "%" in str(non_empty[0]):
            continue

        def _cell(i: Optional[int]) -> Any:
            if i is None or i < 0 or i >= len(cells):
                return None
            return cells[i]

        symbol_raw = _clean_text(_cell(symbol_idx))
        if not symbol_raw:
            for v in cells:
                s = _clean_text(v)
                if not s:
                    continue
                u = s.upper()
                if "USDT" in u and re.search(r"[A-Z]", u.replace("USDT", "")):
                    symbol_raw = s
                    break
                if "/" in u or "PERP" in u or "PERPETUEL" in u or "PERPÉTUEL" in u:
                    symbol_raw = s
                    break

        symbol = _normalize_symbol(symbol_raw)
        ts_cell = _cell(time_idx)
        ts = _to_utc_ts(_clean_text(ts_cell) or ts_cell)
        side_raw = _clean_text(_cell(direction_idx))
        side = None
        if side_raw:
            s = side_raw.strip().lower()
            if s.startswith("acheter"):
                side = "BUY"
            elif s.startswith("vendre"):
                side = "SELL"

        qty_str = _clean_text(_cell(qty_idx))
        exec_price_cell = _cell(exec_price_idx)
        exec_price = _num(_clean_text(exec_price_cell) or exec_price_cell)

        pnl_cell = _cell(pnl_idx)
        if isinstance(pnl_cell, (int, float)):
            pnl_usdt = float(pnl_cell)
        else:
            pnl_usdt = _parse_usdt_amount(_clean_text(pnl_cell)) or _num(_clean_text(pnl_cell) or pnl_cell)

        fee_cell = _cell(fee_idx)
        fee_info = _clean_text(fee_cell)
        status = _clean_text(_cell(status_idx))
        reduce_only_txt = _clean_text(_cell(reduce_only_idx))
        reduce_only = (reduce_only_txt or "").strip().lower() in ("oui", "yes", "true", "1")

        qty = _parse_qty_to_float(qty_str)
        size_usdt = None
        if qty is not None and exec_price is not None:
            size_usdt = qty * exec_price

        fees_candidates = _parse_all_usdt_amounts(fee_info)
        fee_usdt = None
        saved_usdt = None
        if len(fees_candidates) == 1:
            fee_usdt = fees_candidates[0]
        elif len(fees_candidates) >= 2:
            fee_usdt = fees_candidates[0]
            saved_usdt = fees_candidates[1]
        else:
            fee_usdt = _num(_clean_text(fee_cell) or fee_cell)

        if not symbol or ts is None or not side:
            continue

        orders.append(
            {
                "symbol": symbol,
                "time": ts,
                "side": side,
                "reduce_only": reduce_only,
                "qty": qty,
                "exec_price": exec_price,
                "size_usdt": size_usdt,
                "pnl_usdt": pnl_usdt,
                "fee_usdt": fee_usdt,
                "saved_usdt": saved_usdt,
                "status": status,
                "__sheet__": ws.title,
            }
        )

    wb.close()
    return pd.DataFrame(orders)


def mexc_orders_to_trades(df_orders: pd.DataFrame) -> pd.DataFrame:
    if df_orders.empty:
        return df_orders

    df_orders = df_orders.sort_values(by=["time"]).reset_index(drop=True)
    active_positions: Dict[Tuple[str, str], Dict[str, Any]] = {}
    trades: List[Dict[str, Any]] = []

    for _, row in df_orders.iterrows():
        symbol = row.get("symbol")
        side = row.get("side")
        reduce_only = bool(row.get("reduce_only"))

        if side == "BUY":
            inferred_dir = "LONG" if not reduce_only else "SHORT"  # BUY close -> closes SHORT
        else:  # SELL
            inferred_dir = "SHORT" if not reduce_only else "LONG"  # SELL close -> closes LONG

        key = (symbol, inferred_dir)

        qty = _num(row.get("qty"))
        exec_price = _num(row.get("exec_price"))
        size_usdt = _num(row.get("size_usdt"))
        pnl_usdt = _num(row.get("pnl_usdt"))
        fee_usdt = _num(row.get("fee_usdt"))
        saved_usdt = _num(row.get("saved_usdt"))
        ts = _num(row.get("time"))

        if not reduce_only:
            pos = active_positions.get(key)
            if pos is None:
                pos = {
                    "symbol": symbol,
                    "direction": inferred_dir,
                    "open_time": ts,
                    "close_time": None,
                    "entry_qty": 0.0,
                    "entry_cost_usdt": 0.0,
                    "entry_size_usdt": 0.0,
                    "exit_qty": 0.0,
                    "exit_cost_usdt": 0.0,
                    "pnl_usdt": 0.0,
                    "fee_usdt": 0.0,
                    "saved_usdt": 0.0,
                    "last_exit_size_usdt": 0.0,
                    "__sheet__": row.get("__sheet__"),
                }
                active_positions[key] = pos

            if ts is not None:
                if pos.get("open_time") is None or ts < pos.get("open_time"):
                    pos["open_time"] = ts

            if qty is not None:
                pos["entry_qty"] += float(qty)
                if exec_price is not None:
                    pos["entry_cost_usdt"] += float(qty) * float(exec_price)

            if size_usdt is not None:
                pos["entry_size_usdt"] += float(size_usdt)
            continue

        # closing order
        pos = active_positions.get(key)
        if pos is None:
            continue

        if ts is not None:
            if pos.get("close_time") is None or ts > pos.get("close_time"):
                pos["close_time"] = ts

        # Accumuler PnL, frais et économies
        if pnl_usdt is not None:
            pos["pnl_usdt"] += float(pnl_usdt)
        if fee_usdt is not None:
            pos["fee_usdt"] += float(fee_usdt)
        if saved_usdt is not None:
            pos["saved_usdt"] += float(saved_usdt)

        if qty is not None:
            pos["exit_qty"] += float(qty)
            if exec_price is not None:
                pos["exit_cost_usdt"] += float(qty) * float(exec_price)
        
        # Garder la taille de la DERNIÈRE portion fermée pour matcher size_usdt en DB
        # Le bot logue souvent la taille de la portion finale
        if size_usdt is not None:
            pos["last_exit_size_usdt"] = float(size_usdt)

        # Si la position est totalement fermée (ou presque)
        if pos["exit_qty"] >= pos["entry_qty"] - 1e-9:
            entry_price = float(pos["entry_cost_usdt"]) / max(float(pos["entry_qty"]), 1e-12)
            exit_price = float(pos["exit_cost_usdt"]) / max(float(pos["exit_qty"]), 1e-12)
            
            trades.append({
                "symbol": pos["symbol"],
                "direction": pos["direction"],
                "open_time": pos["open_time"],
                "close_time": pos["close_time"],
                "entry_price": entry_price,
                "exit_price": exit_price,
                "size_usdt": pos["last_exit_size_usdt"], # Match DB size_usdt
                "pnl_usdt": pos["pnl_usdt"],
                "fee_usdt": pos["fee_usdt"],
                "saved_usdt": pos["saved_usdt"],
                "__sheet__": pos["__sheet__"]
            })
            del active_positions[key]

    return pd.DataFrame(trades)


@dataclass
class Match:
    exchange_idx: int
    db_id: str
    score: float
    time_diff_s: float
    size_diff_ratio: Optional[float]
    pnl_diff_usdt: Optional[float]


def load_exchange_export(path: str) -> pd.DataFrame:
    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xls"):
        xls = pd.ExcelFile(path)
        # Heuristic: pick best sheet by presence of key columns
        best = None
        for sheet in xls.sheet_names:
            try:
                df = xls.parse(sheet)
            except Exception:
                continue
            if len(df) == 0 or len(df.columns) <= 1:
                continue
            cols_l = {str(c).strip().lower() for c in df.columns}
            score = 0
            for cand in ("symbol", "contract", "pair", "instrument"):
                if cand in cols_l:
                    score += 3
            for cand in ("side", "direction", "position", "type"):
                if cand in cols_l:
                    score += 2
            for cand in ("close_time", "closetime", "timestamp", "time", "close", "closed_time"):
                if cand in cols_l:
                    score += 2
            for cand in ("pnl", "realized_pnl", "realised_pnl", "profit", "pnl_usdt", "profit_usdt"):
                if cand in cols_l:
                    score += 1
            if best is None or score > best[0]:
                best = (score, sheet)

        sheet = best[1] if best else xls.sheet_names[0]
        df = xls.parse(sheet)
        df["__sheet__"] = sheet
        return df
    if ext == ".csv":
        return pd.read_csv(path)
    raise ValueError(f"Unsupported file extension: {ext}")


def load_exchange_export_with_sheet(path: str, sheet: Optional[str]) -> pd.DataFrame:
    ext = os.path.splitext(path)[1].lower()
    if ext not in (".xlsx", ".xls") or not sheet:
        return load_exchange_export(path)
    xls = pd.ExcelFile(path)
    if sheet not in xls.sheet_names:
        raise RuntimeError(f"Sheet '{sheet}' introuvable. Sheets: {xls.sheet_names}")
    df = xls.parse(sheet)
    df["__sheet__"] = sheet
    return df


def load_db_trades(conn) -> List[Dict[str, Any]]:
    cur = conn.cursor(cursor_factory=RealDictCursor)
    cur.execute(
        """
        SELECT
            t.id::text AS id,
            t.symbol,
            t.direction,
            t.timestamp_entry,
            t.timestamp_exit,
            t.entry_price,
            t.exit_price,
            t.size_usdt,
            t.net_pnl_usdt,
            t.net_pnl_pct,
            t.fees_usdt,
            t.exit_reason
        FROM trades t
        WHERE t.timestamp_exit IS NOT NULL
        ORDER BY t.timestamp_exit DESC
        LIMIT 5000
        """
    )
    rows = [dict(r) for r in cur.fetchall()]
    cur.close()
    return rows


def detect_time_offset_seconds(
    df2: pd.DataFrame,
    db: List[Dict[str, Any]],
    max_abs_offset_seconds: int = 86400,
    bin_seconds: int = 60,
) -> Tuple[int, Dict[str, Any]]:
    deltas: List[float] = []
    for _, row in df2.iterrows():
        sym = row.get("_symbol")
        direction = row.get("_direction")
        close_ts = row.get("_close_ts")
        if sym is None or direction is None or close_ts is None:
            continue
        cands = [t for t in db if t.get("_symbol") == sym and t.get("_direction") == direction and t.get("_exit_ts") is not None]
        if not cands:
            continue
        best = min(cands, key=lambda t: abs(float(t["_exit_ts"]) - float(close_ts)))
        delta = float(best["_exit_ts"]) - float(close_ts)
        if abs(delta) <= float(max_abs_offset_seconds):
            deltas.append(delta)

    if not deltas:
        return 0, {"total": 0, "best_bin_count": 0, "best_bin_seconds": None}

    bins: Dict[int, List[float]] = {}
    for d in deltas:
        b = int(round(d / float(bin_seconds))) * int(bin_seconds)
        bins.setdefault(b, []).append(d)

    best_bin = max(bins.items(), key=lambda kv: len(kv[1]))[0]
    best_vals = bins[best_bin]
    best_vals_sorted = sorted(best_vals)
    mid = len(best_vals_sorted) // 2
    if len(best_vals_sorted) % 2 == 1:
        offset = best_vals_sorted[mid]
    else:
        offset = 0.5 * (best_vals_sorted[mid - 1] + best_vals_sorted[mid])

    return int(round(offset)), {"total": len(deltas), "best_bin_count": len(best_vals), "best_bin_seconds": best_bin}


def reconcile(df: pd.DataFrame, db_trades: List[Dict[str, Any]], args) -> None:
    symbol_col = args.col_symbol or _pick_col(df, ["symbol", "contract", "pair", "instrument", "trading_pair"])
    side_col = args.col_side or _pick_col(df, ["side", "direction", "position", "type", "pos_side", "trade_side"])
    close_ts_col = args.col_close_time or _pick_col(
        df,
        [
            "close_time",
            "closetime",
            "closeTime",
            "time",
            "timestamp",
            "close",
            "closed_time",
            "close_timestamp",
            "updated_at",
        ],
    )
    open_ts_col = args.col_open_time or _pick_col(df, ["open_time", "openTime", "open", "opened_time", "open_timestamp", "created_at"])
    pnl_col = args.col_pnl_usdt or _pick_col(df, ["pnl", "realized_pnl", "realised_pnl", "profit", "profit_usdt", "pnl_usdt", "realizedPnl"])
    fee_col = args.col_fee_usdt or _pick_col(df, ["fee", "fees", "commission", "fee_usdt", "commission_usdt", "trading_fee"])
    size_col = args.col_size_usdt or _pick_col(df, ["size_usdt", "notional", "amount_usdt", "value", "qty_usdt", "amount", "qty"])

    missing = [
        ("symbol", symbol_col),
        ("side", side_col),
        ("close_time", close_ts_col),
    ]
    missing = [k for k, v in missing if v is None]
    if missing:
        sample_cols = list(df.columns)[:80]
        raise RuntimeError(
            "Impossible d'auto-détecter les colonnes: "
            + ", ".join(missing)
            + ". Relance avec --col-... (ex: --col-symbol 'Contract').\n"
            + f"Feuille: {df.get('__sheet__', pd.Series(['?'])).iloc[0] if '__sheet__' in df.columns else '?'}\n"
            + f"Colonnes (sample): {sample_cols}"
        )

    df2 = df.copy()
    df2["_symbol"] = df2[symbol_col].map(_str).map(_normalize_symbol)
    df2["_direction"] = df2[side_col].map(_str).map(_normalize_side_to_direction)

    if close_ts_col:
        df2["_close_ts"] = df2[close_ts_col].map(_to_utc_ts)
    else:
        df2["_close_ts"] = None

    if open_ts_col:
        df2["_open_ts"] = df2[open_ts_col].map(_to_utc_ts)
    else:
        df2["_open_ts"] = None

    if pnl_col:
        df2["_pnl_usdt"] = df2[pnl_col].map(_num)
    else:
        df2["_pnl_usdt"] = None

    if fee_col:
        df2["_fee_usdt"] = df2[fee_col].map(_num)
    else:
        df2["_fee_usdt"] = None

    if size_col:
        df2["_size_usdt"] = df2[size_col].map(_num)
    else:
        df2["_size_usdt"] = None

    # Normalize DB
    db = []
    for t in db_trades:
        db.append(
            {
                **t,
                "_symbol": _normalize_symbol(_str(t.get("symbol"))),
                "_direction": _normalize_side_to_direction(_str(t.get("direction"))),
                "_exit_ts": _to_utc_ts(t.get("timestamp_exit")),
                "_entry_ts": _to_utc_ts(t.get("timestamp_entry")),
                "_size_usdt": _num(t.get("size_usdt")),
                "_pnl_usdt": _num(t.get("net_pnl_usdt")),
            }
        )

    if getattr(args, "auto_time_offset", False):
        offset_s, stats = detect_time_offset_seconds(
            df2,
            db,
            max_abs_offset_seconds=getattr(args, "auto_time_offset_max_seconds", 86400),
            bin_seconds=getattr(args, "auto_time_offset_bin_seconds", 60),
        )
        args.time_offset_seconds = offset_s
        print(
            "\nAuto time offset detected (seconds):",
            int(offset_s),
            "| samples:",
            int(stats.get("total", 0)),
            "| best_bin_count:",
            int(stats.get("best_bin_count", 0)),
            "| best_bin_seconds:",
            stats.get("best_bin_seconds"),
        )

    offset_s = int(getattr(args, "time_offset_seconds", 0) or 0)
    if offset_s:
        if "_close_ts" in df2.columns:
            df2["_close_ts"] = df2["_close_ts"].map(lambda v: (v + offset_s) if v is not None and not (isinstance(v, float) and pd.isna(v)) else v)
        if "_open_ts" in df2.columns:
            df2["_open_ts"] = df2["_open_ts"].map(lambda v: (v + offset_s) if v is not None and not (isinstance(v, float) and pd.isna(v)) else v)

    time_tol = args.time_tolerance_seconds

    matches: List[Match] = []
    used_db: set[str] = set()

    for idx, row in df2.iterrows():
        sym = row.get("_symbol")
        direction = row.get("_direction")
        close_ts = row.get("_close_ts")
        open_ts = row.get("_open_ts")

        if sym is None or direction is None:
            continue

        # Candidates: same symbol, and close time near
        candidates = []
        for t in db:
            if t["id"] in used_db:
                continue
            if t.get("_symbol") != sym:
                continue
            if direction and t.get("_direction") and t.get("_direction") != direction:
                continue

            exit_ts = t.get("_exit_ts")
            if close_ts is not None and exit_ts is not None:
                dt = abs(exit_ts - close_ts)
                if dt > time_tol:
                    continue
            elif open_ts is not None and t.get("_entry_ts") is not None:
                dt = abs(t["_entry_ts"] - open_ts)
                if dt > time_tol:
                    continue
            else:
                continue

            # Score
            dt_score = dt
            size_diff_ratio = None
            if row.get("_size_usdt") is not None and t.get("_size_usdt") is not None and t.get("_size_usdt"):
                size_diff_ratio = abs(row["_size_usdt"] - t["_size_usdt"]) / max(t["_size_usdt"], 1e-9)
            size_score = (size_diff_ratio or 0.0) * 1000.0
            score = dt_score + size_score
            candidates.append((score, dt, size_diff_ratio, t))

        if not candidates:
            continue

        candidates.sort(key=lambda x: x[0])
        score, dt, size_diff_ratio, best = candidates[0]
        used_db.add(best["id"])

        pnl_diff_usdt = None
        if row.get("_pnl_usdt") is not None and best.get("_pnl_usdt") is not None:
            pnl_diff_usdt = row["_pnl_usdt"] - best["_pnl_usdt"]

        matches.append(
            Match(
                exchange_idx=int(idx),
                db_id=best["id"],
                score=float(score),
                time_diff_s=float(dt),
                size_diff_ratio=size_diff_ratio,
                pnl_diff_usdt=pnl_diff_usdt,
            )
        )

    # Build report
    matched_df = pd.DataFrame(
        [
            {
                "exchange_idx": m.exchange_idx,
                "db_id": m.db_id,
                "score": m.score,
                "time_diff_s": m.time_diff_s,
                "size_diff_ratio": m.size_diff_ratio,
                "pnl_diff_usdt": m.pnl_diff_usdt,
            }
            for m in matches
        ]
    )

    print("=" * 110)
    print("Reconcile Exchange Export ↔ DB")
    print("export_rows:", len(df2))
    print("db_rows_loaded:", len(db))
    print("matched:", len(matched_df))

    if matched_df.empty:
        print("⚠️ Aucun match trouvé. Augmente --time-tolerance-seconds ou précise les colonnes.")
        return

    # Correlation of pnl when available
    if pnl_col is not None:
        exch_pnl = df2.loc[matched_df["exchange_idx"].values, "_pnl_usdt"].reset_index(drop=True)
        db_pnl = []
        db_by_id = {t["id"]: t for t in db}
        for _i, r in matched_df.iterrows():
            t = db_by_id.get(r["db_id"])
            db_pnl.append(t.get("_pnl_usdt") if t else None)
        db_pnl_s = pd.Series(db_pnl)
        if exch_pnl.notna().sum() >= 5 and db_pnl_s.notna().sum() >= 5:
            corr = pd.concat([exch_pnl, db_pnl_s], axis=1).corr().iloc[0, 1]
            print("\nPnL correlation (exchange vs db):", float(corr) if corr == corr else None)
            print("PnL diff stats (USDT):")
            diffs = (exch_pnl - db_pnl_s).dropna()
            if not diffs.empty:
                print(" mean", float(diffs.mean()), "median", float(diffs.median()), "p95_abs", float(diffs.abs().quantile(0.95)))

    pnl_thr = args.pnl_diff_usdt_threshold
    size_thr = args.size_diff_ratio_threshold
    time_thr = args.time_diff_seconds_threshold

    def _abs(v: Any) -> Optional[float]:
        if v is None or (isinstance(v, float) and pd.isna(v)):
            return None
        return abs(float(v))

    flagged = matched_df.copy()
    flagged["abs_pnl_diff_usdt"] = flagged["pnl_diff_usdt"].map(_abs)
    flagged["abs_size_diff_ratio"] = flagged["size_diff_ratio"].map(_abs)

    problems = flagged[
        (flagged["time_diff_s"] >= time_thr)
        | (flagged["abs_pnl_diff_usdt"].fillna(0) >= pnl_thr)
        | (flagged["abs_size_diff_ratio"].fillna(0) >= size_thr)
    ].sort_values(by=["abs_pnl_diff_usdt", "abs_size_diff_ratio", "time_diff_s"], ascending=False)

    print("\nThresholds:")
    print(" time_diff_s >=", time_thr)
    print(" abs_pnl_diff_usdt >=", pnl_thr)
    print(" abs_size_diff_ratio >=", size_thr)

    print("\nTop mismatches (up to 30):")
    cols = ["exchange_idx", "db_id", "time_diff_s", "pnl_diff_usdt", "size_diff_ratio"]
    print(problems[cols].head(30).to_string(index=False))

    out_csv = args.output_csv
    if out_csv:
        problems.to_csv(out_csv, index=False)
        print("\nSaved:", out_csv)

    print("=" * 110)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True, help="Chemin vers export MEXC (csv/xlsx)")

    ap.add_argument("--sheet", default=None, help="Nom de feuille Excel à utiliser (xlsx/xls uniquement)")
    ap.add_argument("--mexc-fr-orders", action="store_true", help="Parser export MEXC (FR) d'historique d'ordres et reconstruire les trades")

    ap.add_argument("--col-symbol", default=None)
    ap.add_argument("--col-side", default=None)
    ap.add_argument("--col-open-time", default=None)
    ap.add_argument("--col-close-time", default=None)
    ap.add_argument("--col-size-usdt", default=None)
    ap.add_argument("--col-pnl-usdt", default=None)
    ap.add_argument("--col-fee-usdt", default=None)

    ap.add_argument("--time-tolerance-seconds", type=int, default=180)
    ap.add_argument("--time-diff-seconds-threshold", type=int, default=30)
    ap.add_argument("--pnl-diff-usdt-threshold", type=float, default=0.20)
    ap.add_argument("--size-diff-ratio-threshold", type=float, default=0.05)

    ap.add_argument("--time-offset-seconds", type=int, default=0)
    ap.add_argument("--auto-time-offset", action="store_true")
    ap.add_argument("--auto-time-offset-max-seconds", type=int, default=86400)
    ap.add_argument("--auto-time-offset-bin-seconds", type=int, default=60)

    ap.add_argument("--output-csv", default="verification/reconcile_report.csv")

    args = ap.parse_args()

    if args.mexc_fr_orders:
        orders = parse_mexc_fr_orders_xlsx(args.file, args.sheet)
        df = mexc_orders_to_trades(orders)
        # force expected columns
        if args.col_symbol is None:
            args.col_symbol = "symbol"
        if args.col_side is None:
            args.col_side = "direction"
        if args.col_open_time is None:
            args.col_open_time = "open_time"
        if args.col_close_time is None:
            args.col_close_time = "close_time"
        if args.col_size_usdt is None:
            args.col_size_usdt = "size_usdt"
        if args.col_pnl_usdt is None:
            args.col_pnl_usdt = "pnl_usdt"
    else:
        df = load_exchange_export_with_sheet(args.file, args.sheet)

    password = os.environ.get("POSTGRES_PASSWORD") or "@Cmtr1di12345"
    password = quote_plus(password)
    host = os.environ.get("POSTGRES_HOST", "localhost")
    port = os.environ.get("POSTGRES_PORT", "5432")
    dbname = os.environ.get("POSTGRES_DB", "trade_cursor_ml")
    user = os.environ.get("POSTGRES_USER", "postgres")

    conn = psycopg2.connect(f"postgresql://{user}:{password}@{host}:{port}/{dbname}")
    try:
        db_trades = load_db_trades(conn)
    finally:
        conn.close()

    reconcile(df, db_trades, args)


if __name__ == "__main__":
    main()
