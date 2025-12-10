#!/usr/bin/env python3
"""
CHECK EXCEL ORDER FLOW
======================
Vérifie que le fichier Excel exporté contient bien les colonnes order flow.
"""

import sys, os
if sys.platform == 'win32':
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

try:
    import openpyxl
except ImportError:
    print("❌ openpyxl non installé")
    sys.exit(1)

def check_excel_orderflow():
    """Vérifie les colonnes order flow dans le fichier Excel"""
    excel_file = "test_export_orderflow.xlsx"
    
    if not os.path.exists(excel_file):
        print(f"❌ Fichier {excel_file} non trouvé")
        return
    
    print("=" * 60)
    print("  VÉRIFICATION EXCEL ORDER FLOW")
    print("=" * 60)
    
    wb = openpyxl.load_workbook(excel_file)
    
    scan_orderflow = [
        'delta_volume', 'imbalance_normalized', 'spread_volatility_5',
        'book_depth_ratio', 'volume_acceleration', 'price_momentum_5'
    ]
    
    trades_orderflow = [
        'delta_volume', 'imbalance_normalized', 'book_depth_ratio'
    ]
    
    # Vérifier scan_logs
    if 'scan_logs' in wb.sheetnames:
        ws = wb['scan_logs']
        headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
        
        print(f"\n📋 SCAN_LOGS ({len(headers)} colonnes, {ws.max_row-1} lignes)")
        print("🎯 Colonnes Order Flow:")
        
        scan_found = 0
        for col in scan_orderflow:
            if col in headers:
                idx = headers.index(col) + 1
                # Compter les non-null dans cette colonne
                non_null = 0
                for row in range(2, min(52, ws.max_row + 1)):  # 50 premières lignes
                    if ws.cell(row=row, column=idx).value is not None:
                        non_null += 1
                
                print(f"  ✅ {col}: {non_null}/50 non-null")
                scan_found += 1
            else:
                print(f"  ❌ {col}: MANQUANTE")
    
    # Vérifier trades
    if 'trades' in wb.sheetnames:
        ws = wb['trades']
        headers = [cell.value for cell in ws[1]] if ws.max_row > 0 else []
        
        print(f"\n📋 TRADES ({len(headers)} colonnes, {ws.max_row-1} lignes)")
        print("🎯 Colonnes Order Flow:")
        
        trades_found = 0
        for col in trades_orderflow:
            if col in headers:
                idx = headers.index(col) + 1
                # Compter les non-null dans cette colonne
                non_null = 0
                for row in range(2, min(52, ws.max_row + 1)):  # 50 premières lignes
                    if ws.cell(row=row, column=idx).value is not None:
                        non_null += 1
                
                print(f"  ✅ {col}: {non_null}/50 non-null")
                trades_found += 1
            else:
                print(f"  ❌ {col}: MANQUANTE")
    
    print("\n" + "=" * 60)
    print("  RÉSULTAT:")
    print("=" * 60)
    
    if scan_found == 6 and trades_found == 3:
        print("✅ TOUTES les colonnes order flow sont présentes dans l'Excel!")
        print("📊 L'export variablesPanel.exportExcelButton est CORRECT")
    else:
        print(f"❌ Colonnes manquantes: scan_logs {scan_found}/6, trades {trades_found}/3")
        print("🔧 L'export Excel doit être corrigé")
    
    wb.close()

if __name__ == "__main__":
    check_excel_orderflow()
