"""API routes for datalogger exports and maintenance."""
from __future__ import annotations

import json
import logging
from datetime import datetime
from io import BytesIO
from typing import Dict, List

from fastapi import APIRouter
from fastapi.responses import JSONResponse, StreamingResponse

from app.schemas import DataloggerResetResponse
from database.pg import get_cursor

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/api/datalogger", tags=["datalogger"])


@router.get("/export/excel")
async def export_datalogger_excel(start_date: str | None = None, end_date: str | None = None):
    """
    Export datalogger tables to an Excel workbook.
    """
    try:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
        except ImportError:
            return JSONResponse(
                {"error": "openpyxl non installé. Installez-le avec: pip install openpyxl"},
                status_code=500,
            )

        with get_cursor(dict_cursor=True) as cursor:
            wb = Workbook()
            wb.remove(wb.active)

            header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")

            def _normalize_row(row: Dict) -> Dict:
                normalized = {}
                for key, value in row.items():
                    if isinstance(value, (list, dict)):
                        normalized[key] = json.dumps(value, ensure_ascii=False)
                    else:
                        normalized[key] = value
                return normalized

            def _write_sheet(worksheet, rows: List[Dict]):
                if not rows:
                    return
                headers = list(rows[0].keys())
                worksheet.append(headers)
                for cell in worksheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center")
                for row in rows:
                    normalized = _normalize_row(row)
                    worksheet.append([normalized.get(h) for h in headers])
                for col in range(1, len(headers) + 1):
                    worksheet.column_dimensions[get_column_letter(col)].width = 15

            ws_scans = wb.create_sheet("Scans")
            query_scans = """
                SELECT 
                    timestamp, symbol, price, scan_duration_ms,
                    rsi_1m, rsi_5m, score_total,
                    is_opportunity, opportunity_direction, reject_reason,
                    trend_direction, trend_strength
                FROM scan_logs
                WHERE 1=1
            """
            params: List[str] = []
            if start_date:
                query_scans += " AND timestamp >= %s"
                params.append(f"{start_date} 00:00:00")
            if end_date:
                query_scans += " AND timestamp <= %s"
                params.append(f"{end_date} 23:59:59")
            query_scans += " ORDER BY timestamp DESC LIMIT 10000"
            cursor.execute(query_scans, params)
            _write_sheet(ws_scans, cursor.fetchall())

            ws_opps = wb.create_sheet("Opportunities")
            query_opps = """
                SELECT 
                    timestamp, symbol, direction, setup_score,
                    entry_price, tp_price, sl_price,
                    conditions_matched, confirmed_by
                FROM opportunities
                WHERE 1=1
            """
            params_opps: List[str] = []
            if start_date:
                query_opps += " AND timestamp >= %s"
                params_opps.append(f"{start_date} 00:00:00")
            if end_date:
                query_opps += " AND timestamp <= %s"
                params_opps.append(f"{end_date} 23:59:59")
            query_opps += " ORDER BY timestamp DESC LIMIT 10000"
            cursor.execute(query_opps, params_opps)
            _write_sheet(ws_opps, cursor.fetchall())

            ws_trades = wb.create_sheet("Trades")
            query_trades = """
                SELECT 
                    timestamp_entry, timestamp_exit, symbol, direction,
                    entry_price, exit_price, size_usdt,
                    gross_pnl_usdt, net_pnl_usdt, net_pnl_pct,
                    exit_reason, duration_seconds, win
                FROM trades
                WHERE 1=1
            """
            params_trades: List[str] = []
            if start_date:
                query_trades += " AND timestamp_entry >= %s"
                params_trades.append(f"{start_date} 00:00:00")
            if end_date:
                query_trades += " AND timestamp_entry <= %s"
                params_trades.append(f"{end_date} 23:59:59")
            query_trades += " ORDER BY timestamp_entry DESC LIMIT 10000"
            cursor.execute(query_trades, params_trades)
            _write_sheet(ws_trades, cursor.fetchall())

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        filename = (
            f"datalogger_export_{start_date}_{end_date}.xlsx"
            if (start_date and end_date)
            else f"datalogger_export_all_{datetime.now().strftime('%Y%m%d')}.xlsx"
        )
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"},
        )

    except Exception as exc:  # pragma: no cover - unexpected failure path
        logger.exception("Erreur export Excel")
        return JSONResponse({"error": f"Erreur export Excel: {exc}"}, status_code=500)


@router.delete("/reset")
async def reset_datalogger_db():
    """Reset all datalogger tables."""
    try:
        tables = [
            'trades',
            'opportunities',
            'scan_logs',
            'scan_errors',
            'market_context',
            'config_snapshots',
            'trading_sessions',
        ]

        deleted_counts: Dict[str, int] = {}
        with get_cursor() as cursor:
            for table in tables:
                cursor.execute(f"DELETE FROM {table}")
                deleted_counts[table] = cursor.rowcount

        total_deleted = sum(deleted_counts.values())
        logger.warning("🗑️  Base de données PostgreSQL resetée: %s enregistrements supprimés", total_deleted)

        response = DataloggerResetResponse(
            success=True,
            message="Base de données resetée avec succès",
            deleted=deleted_counts,
            total_deleted=total_deleted,
        )
        return JSONResponse(response.model_dump())

    except Exception as exc:  # pragma: no cover - unexpected failure path
        logger.exception("Erreur reset DB")
        return JSONResponse({"error": f"Erreur reset DB: {exc}"}, status_code=500)
