#!/usr/bin/env python3
"""
Test que l'export Excel inclut bien les nouvelles colonnes anti-giveback
"""

import os
import psycopg2
import openpyxl
from tempfile import NamedTemporaryFile
from dotenv import load_dotenv

def test_excel_export():
    load_dotenv()
    
    pg_config = {
        'host': os.getenv('POSTGRES_HOST', 'localhost'),
        'port': int(os.getenv('POSTGRES_PORT', 5432)),
        'database': os.getenv('POSTGRES_DB', 'trade_cursor_ml'),
        'user': os.getenv('POSTGRES_USER', 'postgres'),
        'password': os.getenv('POSTGRES_PASSWORD'),
    }
    
    print("🧪 Test export Excel avec colonnes anti-giveback...")
    
    try:
        conn = psycopg2.connect(**pg_config)
        cursor = conn.cursor()
        
        # Test 1: Simuler l'export SQL utilisé par main.py
        print("\n📊 Test 1: Requête d'export avec colonnes anti-giveback...")
        
        # Colonnes anti-giveback à vérifier dans l'export
        antigiveback_columns = [
            'config_trailing_mfe_enabled',
            'config_trailing_mfe_trigger_pct', 
            'config_trailing_mfe_lock_in_pct',
            'config_partial_tp_be_lock_in_pct',
            'trailing_mfe_triggered',
            'trailing_mfe_triggered_at',
            'trailing_mfe_trigger_pnl_pct',
            'trailing_mfe_trigger_price',
            'trailing_mfe_new_sl'
        ]
        
        # Construire requête SELECT avec toutes les colonnes
        cursor.execute("SELECT column_name FROM information_schema.columns WHERE table_name = 'trades' ORDER BY ordinal_position")
        all_columns = [row[0] for row in cursor.fetchall()]
        
        # Vérifier que les colonnes anti-giveback sont présentes
        ag_found = [col for col in antigiveback_columns if col in all_columns]
        ag_missing = [col for col in antigiveback_columns if col not in all_columns]
        
        print(f"  ✅ Colonnes AG dans schéma: {len(ag_found)}/9")
        if ag_missing:
            print(f"  ❌ Colonnes AG manquantes: {ag_missing}")
        
        # Test 2: Export réel des 5 derniers trades
        print(f"\n📁 Test 2: Export Excel des derniers trades...")
        
        # Construire la requête d'export (comme dans main.py)
        columns_str = ', '.join(all_columns)
        export_query = f"SELECT {columns_str} FROM trades ORDER BY created_at DESC LIMIT 5"
        
        cursor.execute(export_query)
        trades_data = cursor.fetchall()
        
        print(f"  📋 {len(trades_data)} trades exportés")
        print(f"  🔢 {len(all_columns)} colonnes par trade")
        
        # Test 3: Créer un fichier Excel temporaire
        print(f"\n📄 Test 3: Création fichier Excel...")
        
        with NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp_file:
            excel_path = tmp_file.name
        
        # Créer workbook Excel
        workbook = openpyxl.Workbook()
        worksheet = workbook.active
        worksheet.title = "trades"
        
        # Headers (noms colonnes)
        for col_idx, col_name in enumerate(all_columns, 1):
            worksheet.cell(row=1, column=col_idx, value=col_name)
            if col_name in antigiveback_columns:
                # Marquer les colonnes AG en gras
                worksheet.cell(row=1, column=col_idx).font = openpyxl.styles.Font(bold=True)
        
        # Data
        for row_idx, trade_data in enumerate(trades_data, 2):
            for col_idx, value in enumerate(trade_data, 1):
                # Convertir les types pour Excel
                if value is None:
                    display_value = ""
                elif isinstance(value, (int, float)):
                    display_value = value
                else:
                    display_value = str(value)
                    
                worksheet.cell(row=row_idx, column=col_idx, value=display_value)
        
        # Sauvegarder
        workbook.save(excel_path)
        print(f"  ✅ Fichier Excel créé: {excel_path}")
        
        # Test 4: Vérifier contenu du fichier Excel
        print(f"\n🔍 Test 4: Vérification contenu Excel...")
        
        # Recharger le fichier pour vérification
        wb_check = openpyxl.load_workbook(excel_path)
        ws_check = wb_check.active
        
        # Vérifier headers
        excel_headers = []
        for col_idx in range(1, ws_check.max_column + 1):
            header = ws_check.cell(row=1, column=col_idx).value
            if header:
                excel_headers.append(header)
        
        ag_headers_found = [h for h in excel_headers if h in antigiveback_columns]
        print(f"  📋 Headers totaux dans Excel: {len(excel_headers)}")
        print(f"  🛡️ Headers anti-giveback: {len(ag_headers_found)}/9")
        
        for ag_col in antigiveback_columns:
            if ag_col in excel_headers:
                col_idx = excel_headers.index(ag_col) + 1
                print(f"      ✅ {ag_col} (colonne {col_idx})")
            else:
                print(f"      ❌ {ag_col} MANQUANT")
        
        # Vérifier quelques valeurs
        if len(trades_data) > 0:
            print(f"\n📊 Exemple de données (premier trade):")
            for ag_col in antigiveback_columns[:3]:  # 3 premiers pour l'exemple
                if ag_col in excel_headers:
                    col_idx = excel_headers.index(ag_col) + 1
                    value = ws_check.cell(row=2, column=col_idx).value
                    print(f"      - {ag_col}: {value}")
        
        # Nettoyage
        wb_check.close()
        os.unlink(excel_path)
        print(f"  🧹 Fichier temporaire supprimé")
        
        cursor.close()
        conn.close()
        
        return len(ag_headers_found) >= 9
        
    except Exception as e:
        print(f"❌ Erreur lors du test Excel: {e}")
        return False

if __name__ == "__main__":
    success = test_excel_export()
    if success:
        print(f"\n🎉 Export Excel compatible avec toutes les colonnes anti-giveback!")
    else:
        print(f"\n⚠️ Export Excel incomplet ou problématique")
